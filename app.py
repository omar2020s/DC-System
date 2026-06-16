import os
from datetime import datetime
from functools import wraps
from io import BytesIO

from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, Response
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'change-this-secret-key')
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///cement_dc.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(__file__), 'uploads')
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'

# -------------------- Models --------------------
class User(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    full_name = db.Column(db.String(120), nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), default='viewer')  # admin, dc, reviewer, viewer

class Document(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    doc_no = db.Column(db.String(120), unique=True, nullable=False)
    title = db.Column(db.String(250), nullable=False)
    discipline = db.Column(db.String(80), nullable=False)
    doc_type = db.Column(db.String(80), nullable=False)
    revision = db.Column(db.String(30), default='Rev.0')
    status = db.Column(db.String(80), default='Draft')
    workflow_step = db.Column(db.String(80), default='Document Controller')
    submitted_to = db.Column(db.String(120))
    file_name = db.Column(db.String(255))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

class Transmittal(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    transmittal_no = db.Column(db.String(100), unique=True, nullable=False)
    document_id = db.Column(db.Integer, db.ForeignKey('document.id'), nullable=False)
    sent_to = db.Column(db.String(150), nullable=False)
    purpose = db.Column(db.String(100), nullable=False)
    remarks = db.Column(db.Text)
    created_by = db.Column(db.String(120))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    document = db.relationship('Document')

class MAR(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    mar_no = db.Column(db.String(100), unique=True, nullable=False)
    material = db.Column(db.String(200), nullable=False)
    supplier = db.Column(db.String(150))
    discipline = db.Column(db.String(80), default='Civil')
    status = db.Column(db.String(80), default='Submitted')
    remarks = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class RFI(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rfi_no = db.Column(db.String(100), unique=True, nullable=False)
    subject = db.Column(db.String(250), nullable=False)
    location = db.Column(db.String(150))
    question = db.Column(db.Text)
    answer = db.Column(db.Text)
    status = db.Column(db.String(80), default='Open')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class NCR(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ncr_no = db.Column(db.String(100), unique=True, nullable=False)
    description = db.Column(db.Text, nullable=False)
    location = db.Column(db.String(150))
    root_cause = db.Column(db.Text)
    corrective_action = db.Column(db.Text)
    status = db.Column(db.String(80), default='Open')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

def roles_required(*roles):
    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if current_user.role not in roles and current_user.role != 'admin':
                flash('You are not authorized for this action.', 'danger')
                return redirect(url_for('dashboard'))
            return fn(*args, **kwargs)
        return wrapper
    return decorator

def save_file(file):
    if file and file.filename:
        name = datetime.now().strftime('%Y%m%d%H%M%S_') + secure_filename(file.filename)
        file.save(os.path.join(app.config['UPLOAD_FOLDER'], name))
        return name
    return None

# -------------------- Auth --------------------
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        user = User.query.filter_by(username=request.form['username']).first()
        if user and check_password_hash(user.password_hash, request.form['password']):
            login_user(user)
            return redirect(url_for('dashboard'))
        flash('Wrong username or password.', 'danger')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

# -------------------- Dashboard --------------------
@app.route('/')
@login_required
def dashboard():
    data = {
        'documents': Document.query.count(),
        'approved': Document.query.filter_by(status='Approved').count(),
        'review': Document.query.filter_by(status='Under Review').count(),
        'rejected': Document.query.filter_by(status='Rejected').count(),
        'mar': MAR.query.count(),
        'rfi': RFI.query.count(),
        'ncr': NCR.query.count(),
    }
    return render_template('dashboard.html', data=data)

# -------------------- Documents --------------------
@app.route('/documents')
@login_required
def documents():
    q = request.args.get('q', '')
    query = Document.query
    if q:
        query = query.filter((Document.doc_no.contains(q)) | (Document.title.contains(q)) | (Document.discipline.contains(q)))
    docs = query.order_by(Document.updated_at.desc()).all()
    return render_template('documents.html', docs=docs, q=q)

@app.route('/documents/add', methods=['GET', 'POST'])
@login_required
@roles_required('admin', 'dc')
def add_document():
    if request.method == 'POST':
        doc = Document(
            doc_no=request.form['doc_no'], title=request.form['title'], discipline=request.form['discipline'],
            doc_type=request.form['doc_type'], revision=request.form['revision'], status=request.form['status'],
            workflow_step=request.form['workflow_step'], submitted_to=request.form.get('submitted_to'),
            file_name=save_file(request.files.get('file'))
        )
        db.session.add(doc); db.session.commit()
        flash('Document added successfully.', 'success')
        return redirect(url_for('documents'))
    return render_template('document_form.html', doc=None)

@app.route('/documents/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@roles_required('admin', 'dc', 'reviewer')
def edit_document(id):
    doc = Document.query.get_or_404(id)
    if request.method == 'POST':
        doc.title = request.form['title']; doc.discipline = request.form['discipline']; doc.doc_type = request.form['doc_type']
        doc.revision = request.form['revision']; doc.status = request.form['status']; doc.workflow_step = request.form['workflow_step']
        doc.submitted_to = request.form.get('submitted_to')
        new_file = save_file(request.files.get('file'))
        if new_file: doc.file_name = new_file
        db.session.commit(); flash('Document updated successfully.', 'success')
        return redirect(url_for('documents'))
    return render_template('document_form.html', doc=doc)

@app.route('/uploads/<path:filename>')
@login_required
def uploads(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/documents/export')
@login_required
def export_documents():
    return export_model(Document.query.all(), 'Document_Register', ['doc_no','title','discipline','doc_type','revision','status','workflow_step','submitted_to','created_at'])

# -------------------- Workflow --------------------
@app.route('/workflow/<int:id>/<action>')
@login_required
@roles_required('admin', 'reviewer', 'dc')
def workflow_action(id, action):
    doc = Document.query.get_or_404(id)
    if action == 'submit':
        doc.status = 'Under Review'; doc.workflow_step = 'Consultant Review'
    elif action == 'approve':
        doc.status = 'Approved'; doc.workflow_step = 'Closed'
    elif action == 'comment':
        doc.status = 'Approved with Comments'; doc.workflow_step = 'Contractor Revision'
    elif action == 'reject':
        doc.status = 'Rejected'; doc.workflow_step = 'Contractor Revision'
    db.session.commit()
    flash('Workflow updated.', 'success')
    return redirect(url_for('documents'))

# -------------------- Transmittals --------------------
@app.route('/transmittals')
@login_required
def transmittals():
    items = Transmittal.query.order_by(Transmittal.created_at.desc()).all()
    docs = Document.query.order_by(Document.doc_no).all()
    return render_template('transmittals.html', items=items, docs=docs)

@app.route('/transmittals/add', methods=['POST'])
@login_required
@roles_required('admin', 'dc')
def add_transmittal():
    t = Transmittal(
        transmittal_no=request.form['transmittal_no'], document_id=request.form['document_id'], sent_to=request.form['sent_to'],
        purpose=request.form['purpose'], remarks=request.form.get('remarks'), created_by=current_user.full_name
    )
    db.session.add(t); db.session.commit(); flash('Transmittal created.', 'success')
    return redirect(url_for('transmittals'))

@app.route('/transmittals/<int:id>/pdf')
@login_required
def transmittal_pdf(id):
    t = Transmittal.query.get_or_404(id)
    buffer = BytesIO(); p = canvas.Canvas(buffer, pagesize=A4); w, h = A4
    p.setFont('Helvetica-Bold', 16); p.drawString(2*cm, h-2*cm, 'TRANSMITTAL FORM')
    p.setFont('Helvetica', 10)
    lines = [
        ('Project', 'Cement Plant Construction Project'), ('Transmittal No', t.transmittal_no),
        ('Date', t.created_at.strftime('%Y-%m-%d')), ('Sent To', t.sent_to), ('Purpose', t.purpose),
        ('Document No', t.document.doc_no), ('Document Title', t.document.title), ('Revision', t.document.revision),
        ('Status', t.document.status), ('Remarks', t.remarks or '-'), ('Created By', t.created_by or '-')
    ]
    y = h-3.5*cm
    for key, val in lines:
        p.setFont('Helvetica-Bold', 10); p.drawString(2*cm, y, f'{key}:')
        p.setFont('Helvetica', 10); p.drawString(6*cm, y, str(val)[:90]); y -= 0.7*cm
    p.line(2*cm, 4*cm, 18*cm, 4*cm)
    p.drawString(2*cm, 3.4*cm, 'Received By: ____________________')
    p.drawString(10*cm, 3.4*cm, 'Signature: ____________________')
    p.save(); buffer.seek(0)
    return Response(buffer, mimetype='application/pdf', headers={'Content-Disposition': f'inline; filename={t.transmittal_no}.pdf'})

# -------------------- Logs Generic --------------------
def export_model(items, filename, fields):
    wb = Workbook(); ws = wb.active; ws.title = filename[:31]
    ws.append(fields)
    for item in items:
        row = []
        for f in fields:
            v = getattr(item, f)
            row.append(v.strftime('%Y-%m-%d') if isinstance(v, datetime) else v)
        ws.append(row)
    stream = BytesIO(); wb.save(stream); stream.seek(0)
    return Response(stream, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename={filename}.xlsx'})

@app.route('/mar', methods=['GET', 'POST'])
@login_required
def mar_log():
    if request.method == 'POST':
        db.session.add(MAR(mar_no=request.form['mar_no'], material=request.form['material'], supplier=request.form.get('supplier'), discipline=request.form.get('discipline'), status=request.form.get('status'), remarks=request.form.get('remarks')))
        db.session.commit(); flash('MAR added.', 'success'); return redirect(url_for('mar_log'))
    return render_template('mar.html', items=MAR.query.order_by(MAR.created_at.desc()).all())

@app.route('/rfi', methods=['GET', 'POST'])
@login_required
def rfi_log():
    if request.method == 'POST':
        db.session.add(RFI(rfi_no=request.form['rfi_no'], subject=request.form['subject'], location=request.form.get('location'), question=request.form.get('question'), answer=request.form.get('answer'), status=request.form.get('status')))
        db.session.commit(); flash('RFI added.', 'success'); return redirect(url_for('rfi_log'))
    return render_template('rfi.html', items=RFI.query.order_by(RFI.created_at.desc()).all())

@app.route('/ncr', methods=['GET', 'POST'])
@login_required
def ncr_log():
    if request.method == 'POST':
        db.session.add(NCR(ncr_no=request.form['ncr_no'], description=request.form['description'], location=request.form.get('location'), root_cause=request.form.get('root_cause'), corrective_action=request.form.get('corrective_action'), status=request.form.get('status')))
        db.session.commit(); flash('NCR added.', 'success'); return redirect(url_for('ncr_log'))
    return render_template('ncr.html', items=NCR.query.order_by(NCR.created_at.desc()).all())

@app.route('/export/<name>')
@login_required
def export_log(name):
    if name == 'mar': return export_model(MAR.query.all(), 'MAR_Log', ['mar_no','material','supplier','discipline','status','remarks','created_at'])
    if name == 'rfi': return export_model(RFI.query.all(), 'RFI_Log', ['rfi_no','subject','location','question','answer','status','created_at'])
    if name == 'ncr': return export_model(NCR.query.all(), 'NCR_Log', ['ncr_no','description','location','root_cause','corrective_action','status','created_at'])
    return redirect(url_for('dashboard'))

# -------------------- Users --------------------
@app.route('/users', methods=['GET', 'POST'])
@login_required
@roles_required('admin')
def users():
    if request.method == 'POST':
        u = User(username=request.form['username'], full_name=request.form['full_name'], password_hash=generate_password_hash(request.form['password']), role=request.form['role'])
        db.session.add(u); db.session.commit(); flash('User created.', 'success'); return redirect(url_for('users'))
    return render_template('users.html', users=User.query.all())

# -------------------- Init --------------------
def seed():
    if not User.query.filter_by(username='admin').first():
        db.session.add(User(username='admin', full_name='System Admin', password_hash=generate_password_hash('admin123'), role='admin'))
        db.session.add(User(username='dc', full_name='Document Controller', password_hash=generate_password_hash('dc123'), role='dc'))
        db.session.add(User(username='reviewer', full_name='Consultant Reviewer', password_hash=generate_password_hash('review123'), role='reviewer'))
        db.session.commit()

with app.app_context():
    db.create_all(); seed()

if __name__ == '__main__':
    app.run(debug=True)
